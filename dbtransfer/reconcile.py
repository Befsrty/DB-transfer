"""Сверка активных данных ADSO с эталоном (ТЗ, ФП-8).

Сравнение — чистая функция compare() над двумя DataFrame, без базы: её проверяют модульные тесты.
Колонки эталона называются по техническим именам InfoObject; другое соответствие задаётся
в reconcile.yaml проекта.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
import psycopg
import yaml

from dbtransfer import naming
from dbtransfer.analysis.excel import _table, _title
from dbtransfer.generator.tables import TableBuilder
from dbtransfer.generator.types import TypeRules, is_key_figure
from dbtransfer.model import Adso, Model, ObjectType

MATCH, DIFF, NOT_IN_DB, NOT_IN_ETALON = "совпало", "расхождение", "нет в ADB", "нет в эталоне"
STATUSES = (MATCH, DIFF, NOT_IN_DB, NOT_IN_ETALON)
DEFAULT_TOLERANCE = 0.01
RECONCILE_FILE = "reconcile.yaml"


class ReconcileError(Exception):
    """Сверка не выполнена; текст понятен пользователю."""


@dataclass
class Spec:
    """Что сверяем: ключ, показатели и их единицы (валюта или единица измерения)."""

    object_name: str
    keys: list[str]       # InfoObject ключа
    measures: list[str]   # InfoObject показателей
    units: dict[str, str]  # показатель -> InfoObject единицы, если он есть в объекте
    columns: dict[str, str]  # InfoObject -> колонка в базе
    datatypes: dict[str, str]


@dataclass
class ReconcileResult:
    object_name: str
    etalon_file: str
    detail: pd.DataFrame   # строка на ключ: ключ, по каждому показателю эталон/база/разница, статус
    totals: pd.DataFrame   # суммы показателей по единицам
    spec: Spec
    run_id: int | None = None
    excel: Path | None = None
    lower: list[ReconcileResult] = field(default_factory=list)

    def count(self, status: str) -> int:
        return int((self.detail["status"] == status).sum()) if len(self.detail) else 0

    @property
    def ok(self) -> bool:
        return len(self.detail) > 0 and self.count(MATCH) == len(self.detail)

    @property
    def problems(self) -> pd.DataFrame:
        return self.detail[self.detail["status"] != MATCH]


def make_spec(model: Model, object_name: str) -> Spec:
    adso = model.adsos.get(object_name)
    if adso is None:
        raise ReconcileError(f"ADSO {object_name} не найден. Сверяются активные данные ADSO, например ZSD_A02")
    builder = TableBuilder(model, TypeRules.load())
    keys = [f.name for f in builder.adso_keys(adso)] or [f.name for f in adso.fields if not is_key_figure(model, f)]
    measures = [f.name for f in adso.fields if is_key_figure(model, f)]
    names = {f.name for f in adso.fields}
    units = {}
    for m in measures:
        iobj = model.info_objects.get(m)
        if iobj and iobj.unit_info_object in names:
            units[m] = iobj.unit_info_object
    return Spec(object_name, keys, measures, units,
                {f.name: naming.column(f.name, f.info_object) for f in adso.fields},
                {f.name: f.datatype for f in adso.fields})


def _text(value, datatype: str) -> str:
    """Значение признака в виде, как в эталоне: дата ГГГГММДД, пусто для NULL."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, date | datetime):
        return value.strftime("%Y%m%d")
    text = str(value)
    if datatype == "DATS" and text in ("00000000",):
        return ""
    return text


def _number(value) -> float | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, str):
        value = value.strip().replace(" ", "").replace(",", ".")
        return float(value) if value else None
    return float(value)


def read_etalon(path: Path, spec: Spec, columns: dict[str, str] | None = None) -> pd.DataFrame:
    """Эталон CSV -> DataFrame с колонками InfoObject (ключи — строки, показатели — числа)."""
    if not path.is_file():
        raise ReconcileError(f"Файл эталона {path} не найден")
    df = pd.read_csv(path, sep=None, engine="python", dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df = df.rename(columns={k: v for k, v in (columns or {}).items()})
    missing = [c for c in spec.keys + spec.measures if c not in df.columns]
    if missing:
        raise ReconcileError(f"В эталоне {path.name} нет колонок {', '.join(missing)}. Колонки называются "
                             f"по InfoObject; другое соответствие задаётся в {RECONCILE_FILE} проекта")
    result = pd.DataFrame({k: [_text(v, spec.datatypes.get(k, "")) for v in df[k]] for k in spec.keys})
    for m in spec.measures:
        try:
            result[m] = [_number(v) for v in df[m]]
        except ValueError as e:
            raise ReconcileError(f"В эталоне {path.name} в колонке {m} не число: {e}") from e
    for k in set(spec.units.values()) - set(spec.keys):
        result[k] = [str(v) for v in df[k]] if k in df.columns else ""
    return result


def read_actual(conn: psycopg.Connection, model: Model, spec: Spec) -> pd.DataFrame:
    table = naming.data_table(model, model.adsos[spec.object_name])
    names = list(dict.fromkeys(spec.keys + spec.measures + list(spec.units.values())))
    try:
        with conn.cursor() as cur:
            rows = cur.execute(f"SELECT {', '.join(spec.columns[n] for n in names)} FROM {table}").fetchall()
        conn.rollback()
    except psycopg.Error as e:
        conn.rollback()
        raise ReconcileError(f"Не удалось прочитать {table}: {(e.diag.message_primary or str(e)).strip()}. "
                             "Разверните пакет и выполните прогон: dbtransfer deploy, load, run") from e
    data = {n: [r[i] for r in rows] for i, n in enumerate(names)}
    result = pd.DataFrame({k: [_text(v, spec.datatypes.get(k, "")) for v in data[k]] for k in spec.keys})
    for m in spec.measures:
        result[m] = [_number(v) if not isinstance(v, Decimal) else float(v) for v in data[m]]
    for k in set(spec.units.values()) - set(spec.keys):
        result[k] = [_text(v, "") for v in data[k]]
    return result


def compare(etalon: pd.DataFrame, actual: pd.DataFrame, spec: Spec,
            tolerance: float = DEFAULT_TOLERANCE) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Сверка по ключу и каждому показателю с допуском (ФП-8.2). Возвращает построчный итог и суммы
    по показателям в разрезе единиц."""
    for name, df in (("эталоне", etalon), ("базе", actual)):
        dup = df[df.duplicated(spec.keys, keep=False)]
        if len(dup):
            raise ReconcileError(f"В {name} повторяется ключ {', '.join(spec.keys)}: "
                                 f"{' / '.join(dup.iloc[0][spec.keys])}")
    merged = etalon.merge(actual, on=spec.keys, how="outer", suffixes=("_bw", "_adb"), indicator=True)
    detail = merged[spec.keys].copy()
    for m in spec.measures:
        bw, adb = merged[f"{m}_bw"], merged[f"{m}_adb"]
        detail[f"{m}_bw"] = bw
        detail[f"{m}_adb"] = adb
        detail[f"{m}_diff"] = (adb.fillna(0) - bw.fillna(0)).round(6)
    within = pd.Series(True, index=merged.index)
    for m in spec.measures:
        within &= detail[f"{m}_diff"].abs() <= tolerance + 1e-9
    detail["status"] = [
        NOT_IN_DB if side == "left_only" else NOT_IN_ETALON if side == "right_only" else MATCH if ok else DIFF
        for side, ok in zip(merged["_merge"], within, strict=True)]
    detail = detail.sort_values(spec.keys).reset_index(drop=True)

    totals = []
    for m in spec.measures:
        unit = spec.units.get(m)
        if unit:
            bw_units, adb_units = etalon[unit], actual[unit]
            groups = sorted(set(bw_units) | set(adb_units))
            for u in groups:
                b = etalon.loc[bw_units == u, m].fillna(0).sum()
                a = actual.loc[adb_units == u, m].fillna(0).sum()
                totals.append([m, u, round(b, 6), round(a, 6), round(a - b, 6)])
        else:
            b, a = etalon[m].fillna(0).sum(), actual[m].fillna(0).sum()
            totals.append([m, "", round(b, 6), round(a, 6), round(a - b, 6)])
    return detail, pd.DataFrame(totals, columns=["measure", "unit", "bw", "adb", "diff"])


def load_columns(project_dir: Path | None, object_name: str) -> tuple[dict[str, str], float]:
    """Соответствие колонок эталона и допуск из reconcile.yaml проекта."""
    path = project_dir / RECONCILE_FILE if project_dir else None
    if path is None or not path.is_file():
        return {}, DEFAULT_TOLERANCE
    try:
        raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError as e:
        raise ReconcileError(f"{path}: файл не читается: {e}") from e
    spec = raw.get(object_name) or {}
    return dict(spec.get("columns") or {}), float(spec.get("tolerance", raw.get("tolerance", DEFAULT_TOLERANCE)))


def save_result(conn: psycopg.Connection, model: Model, result: ReconcileResult) -> int:
    """Протокол в meta.reconcile_result: строка на ключ и показатель (ФП-8.3)."""
    meta = naming.schema("meta", model.schema_prefix)
    spec = result.spec
    rows = []
    for _, r in result.detail.iterrows():
        key = json.dumps({k: r[k] for k in spec.keys}, ensure_ascii=False)
        for m in spec.measures:
            rows.append((result.object_name, result.etalon_file, key, m, _none(r[f"{m}_bw"]), _none(r[f"{m}_adb"]),
                         _none(r[f"{m}_diff"]), r["status"]))
    with conn.cursor() as cur:
        run_id = cur.execute(f"SELECT nextval('{meta}.reconcile_run_seq')").fetchone()[0]
        cur.executemany(
            f"INSERT INTO {meta}.reconcile_result (run_id, object_name, etalon_file, key_values, measure, value_bw, "
            "value_adb, diff, status, run_at) VALUES (%s, %s, %s, %s::jsonb, %s, %s, %s, %s, %s, now())",
            [(run_id, *row) for row in rows])
    conn.commit()
    return run_id


def _none(value):
    return None if value is None or pd.isna(value) else float(value)


def write_excel(result: ReconcileResult, path: Path) -> Path:
    from openpyxl import Workbook

    spec = result.spec
    wb = Workbook()
    ws = wb.active
    ws.title = "Итог"
    row = _title(ws, 1, f"Сверка {result.object_name} с эталоном {result.etalon_file}")
    row += 1
    row = _title(ws, row, "Строки по статусам")
    row = _table(ws, ["Статус", "Строк"], [[s, result.count(s)] for s in STATUSES] + [["Всего", len(result.detail)]],
                 start_row=row, widths=[36, 18, 18, 18, 18])
    row = _title(ws, row, "Суммы по показателям")
    row = _table(ws, ["Показатель", "Единица", "Эталон", "ADB", "Разница"], result.totals.values.tolist(),
                 start_row=row)
    for lower in result.lower:
        verdict = "совпал" if lower.ok else f"не совпал: {len(lower.problems)} строк с отличиями"
        ws.cell(row=row, column=1, value=f"Нижний слой {lower.object_name}: {verdict}")
        row += 1

    ws = wb.create_sheet("Расхождения")
    header = [*spec.keys]
    for m in spec.measures:
        header += [f"{m} эталон", f"{m} ADB", f"{m} разница"]
    header.append("Статус")
    rows = []
    for _, r in result.problems.iterrows():
        values = [r[k] for k in spec.keys]
        for m in spec.measures:
            values += [_none(r[f"{m}_bw"]), _none(r[f"{m}_adb"]), _none(r[f"{m}_diff"])]
        rows.append([*values, r["status"]])
    _table(ws, header, rows or [["Расхождений нет"]], widths=[16] * len(header))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def upstream_adsos(model: Model, object_name: str) -> list[str]:
    return sorted({t.source.name for t in model.transformations.values()
                   if t.target.name == object_name and t.source.type == ObjectType.ADSO})


def reconcile(conn: psycopg.Connection, model: Model, object_name: str, etalon: Path, out_dir: Path,
              project_dir: Path | None = None, tolerance: float | None = None, check_lower: bool = True,
              now: datetime | None = None) -> ReconcileResult:
    """Сверить объект с эталоном, сохранить протокол в базе и Excel (ФП-8.1–8.4)."""
    spec = make_spec(model, object_name)
    columns, default_tol = load_columns(project_dir, object_name)
    tol = default_tol if tolerance is None else tolerance
    detail, totals = compare(read_etalon(etalon, spec, columns), read_actual(conn, model, spec), spec, tol)
    result = ReconcileResult(object_name, etalon.name, detail, totals, spec)
    # ФП-8.4: при расхождении проверяем нижний слой, если рядом лежит его эталон
    if check_lower and not result.ok:
        for lower in upstream_adsos(model, object_name):
            path = etalon.with_name(f"etalon_{lower}.csv")
            if path.is_file() and isinstance(model.adsos.get(lower), Adso):
                result.lower.append(reconcile(conn, model, lower, path, out_dir, project_dir, tolerance,
                                              check_lower=False, now=now))
    result.run_id = save_result(conn, model, result)
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    result.excel = write_excel(result, out_dir / f"reconcile_{object_name}_{stamp}.xlsx")
    return result
