"""Отчёты в Excel: проверка файлов выгрузки (ФП-2.3) и отчёт о сложности (ФП-3.4, ФП-3.7)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dbtransfer.analysis.assessment import TYPE_ORDER, Report
from dbtransfer.config import load_bw_tables
from dbtransfer.model import Model, Severity

HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")
LEVEL_FILL = {1: "E3F2E1", 2: "FFF4D6", 3: "FBE0DD"}
SEVERITY_LABELS = {Severity.ERROR: "ошибка", Severity.WARNING: "предупреждение", Severity.INFO: "сведения"}


def _table(ws: Worksheet, header: list[str], rows: list[list], start_row: int = 1,
           widths: list[int] | None = None) -> int:
    """Записать таблицу с заголовком; вернуть номер строки после неё."""
    for col, title in enumerate(header, start=1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for r, values in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value).alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    return start_row + len(rows) + 2


def _title(ws: Worksheet, row: int, text: str) -> int:
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=12)
    return row + 1


def write_import_check(model: Model, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Файлы"
    specs = load_bw_tables()
    rows = []
    for name, spec in specs.items():
        info = model.files.get(name, {})
        rows.append([name, spec.group, "да" if spec.required else "нет", info.get("file", ""),
                     info.get("status", ""), info.get("encoding", ""), info.get("delimiter", ""),
                     info.get("rows", 0), "да" if spec.hypothesis else ""])
    _table(ws, ["Таблица", "Группа", "Обязательная", "Файл", "Статус", "Кодировка", "Разделитель", "Строк",
                "Состав полей — гипотеза"], rows, widths=[14, 22, 13, 18, 12, 14, 12, 8, 14])

    ws = wb.create_sheet("Проблемы")
    order = {Severity.ERROR: 0, Severity.WARNING: 1, Severity.INFO: 2}
    issues = sorted(model.issues, key=lambda i: (order[i.severity], i.file or "", i.line or 0))
    _table(ws, ["Серьёзность", "Файл", "Строка", "Объект", "Сообщение"],
           [[SEVERITY_LABELS[i.severity], i.file or "", i.line, i.object or "", i.message] for i in issues],
           widths=[16, 18, 8, 34, 100])

    ws = wb.create_sheet("Неактивные")
    _table(ws, ["Тип", "Объект", "Версии", "Файл"],
           [[o.type, o.name, ", ".join(o.versions), o.file] for o in model.inactive], widths=[20, 30, 10, 18])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_complexity(model: Model, report: Report, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Итог"
    row = _title(ws, 1, f"Отчёт о сложности переноса: {model.name or model.source}")
    ws.cell(row=row, column=1, value=f"Импорт {model.imported_at:%d.%m.%Y %H:%M}" if model.imported_at else "")
    row += 2

    row = _title(ws, row, "Объекты по типам и уровням автоматизации")
    counts = []
    levels_by_type: dict[str, list[int]] = {}
    for a in report.objects + report.unparsed:
        levels_by_type.setdefault(a.type, [0, 0, 0])[a.level - 1] += 1
    for t in TYPE_ORDER + sorted(set(levels_by_type) - set(TYPE_ORDER)):
        if t in levels_by_type:
            lv = levels_by_type[t]
            counts.append([t, sum(lv), *lv])
    counts.append(["Всего", report.total, report.by_level[1], report.by_level[2], report.by_level[3]])
    row = _table(ws, ["Тип", "Всего", "Уровень 1", "Уровень 2", "Уровень 3"], counts, start_row=row,
                 widths=[32, 10, 12, 12, 12])
    ws.cell(row=row - 2, column=1).font = Font(bold=True)

    row = _title(ws, row, "Метрики эффекта")
    metrics = [
        ["Объектов уровня 1 (переводятся автоматически)", f"{report.by_level[1]} из {report.total}",
         f"{report.level1_share:.0%}"],
        ["Ручная работа по нормам effort_norms.yaml, часов", f"{report.hours:g}", ""],
        ["Сгенерировано строк SQL и объектов БД", "появится в шаге 3 (генерация SQL)", ""],
    ]
    row = _table(ws, ["Показатель", "Значение", "Доля"], metrics, start_row=row)

    row = _title(ws, row, "Flat File DataSource")
    _table(ws, ["DataSource"], [[f] for f in report.flat_files] or [["нет"]], start_row=row)

    ws = wb.create_sheet("Объекты")
    rows = [[a.type, a.name, a.description, a.level, a.reason, a.hours or None] for a in report.objects]
    _table(ws, ["Тип", "Объект", "Описание", "Уровень", "Причина", "Часов"], rows,
           widths=[28, 34, 40, 9, 70, 8])
    for r, a in enumerate(report.objects, start=2):
        ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor=LEVEL_FILL[a.level])
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Рутины")
    _table(ws, ["Трансформация", "Источник", "Цель", "Вид", "CODEID", "Строк кода", "Код"],
           [[r.transformation, r.source, r.target, r.kind, r.code_id, r.lines, r.code or "кода нет в выгрузке"]
            for r in report.routines], widths=[36, 24, 14, 26, 36, 11, 90])

    ws = wb.create_sheet("Не разобрано")
    _table(ws, ["Объект", "Где", "Причина", "Часов"],
           [[a.name, a.description, a.reason, a.hours] for a in report.unparsed] or [["нет", "", "", None]],
           widths=[34, 24, 100, 8])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
