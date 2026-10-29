"""Сверка на подготовленных DataFrame (ТЗ 4.7, ФП-8)."""

import pandas as pd
import pytest
from conftest import DEMO
from openpyxl import load_workbook

from dbtransfer.importers.bw_csv import import_bw_export
from dbtransfer.reconcile import (
    DIFF,
    MATCH,
    NOT_IN_DB,
    NOT_IN_ETALON,
    ReconcileError,
    ReconcileResult,
    compare,
    load_columns,
    make_spec,
    read_etalon,
    upstream_adsos,
    write_excel,
)


@pytest.fixture(scope="module")
def model():
    return import_bw_export(DEMO / "bw_export")


@pytest.fixture(scope="module")
def spec(model):
    return make_spec(model, "ZSD_A02")


def frame(rows):
    cols = ["0SOLD_TO", "0CALMONTH", "0DOC_CURRCY", "0SALES_UNIT", "0NETVAL_INV", "0INV_QTY"]
    return pd.DataFrame(rows, columns=cols)


def test_spec(spec, model):
    assert spec.keys == ["0SOLD_TO", "0CALMONTH", "0DOC_CURRCY", "0SALES_UNIT"]
    assert spec.measures == ["0NETVAL_INV", "0INV_QTY"]
    assert spec.units == {"0NETVAL_INV": "0DOC_CURRCY", "0INV_QTY": "0SALES_UNIT"}
    a01 = make_spec(model, "ZSD_A01")
    assert a01.keys == ["0DOC_NUMBER", "0S_ORD_ITEM"] and a01.columns["0CALDAY"] == "calday"
    with pytest.raises(ReconcileError, match="не найден"):
        make_spec(model, "ZSALES_FF")


def test_compare_statuses(spec):
    etalon = frame([["1", "202609", "RUB", "ST", 100.0, 1.0], ["1", "202610", "RUB", "ST", 50.0, 2.0],
                    ["2", "202609", "EUR", "ST", 10.0, 1.0], ["3", "202609", "RUB", "ST", 0.0, 0.0]])
    actual = frame([["1", "202609", "RUB", "ST", 100.004, 1.0], ["1", "202610", "RUB", "ST", 49.0, 2.0],
                    ["4", "202609", "RUB", "ST", 5.0, 1.0], ["3", "202609", "RUB", "ST", 0.0, 0.0]])
    detail, totals = compare(etalon, actual, spec)
    status = {r["0SOLD_TO"] + r["0CALMONTH"]: r["status"] for _, r in detail.iterrows()}
    assert status == {"1202609": MATCH, "1202610": DIFF, "2202609": NOT_IN_DB, "4202609": NOT_IN_ETALON,
                      "3202609": MATCH}
    row = detail[detail["status"] == DIFF].iloc[0]
    assert row["0NETVAL_INV_diff"] == -1.0 and row["0INV_QTY_diff"] == 0
    rub = totals[(totals.measure == "0NETVAL_INV") & (totals.unit == "RUB")].iloc[0]
    assert (rub.bw, rub.adb) == (150.0, 154.004)
    assert set(totals[totals.measure == "0NETVAL_INV"].unit) == {"EUR", "RUB"}
    # Допуск: 0,01 по умолчанию, можно задать больше
    detail, _ = compare(etalon, actual, spec, tolerance=2)
    assert (detail["status"] == DIFF).sum() == 0


def test_duplicate_key_is_error(spec):
    rows = [["1", "202609", "RUB", "ST", 1.0, 1.0]] * 2
    with pytest.raises(ReconcileError, match="повторяется ключ"):
        compare(frame(rows), frame(rows[:1]), spec)


def test_read_etalon(spec, model, tmp_path):
    df = read_etalon(DEMO / "data" / "etalon_ZSD_A02.csv", spec)
    assert len(df) == 142 and df["0NETVAL_INV"].dtype == float
    a01 = make_spec(model, "ZSD_A01")
    df = read_etalon(DEMO / "data" / "micro" / "expected_ZSD_A01.csv", a01)
    assert list(df["0NETVAL_INV"]) == [800.0, 500.0, 250.0, 300.0, 100.0, 400.0]

    (tmp_path / "e.csv").write_text("Клиент;0CALMONTH;0DOC_CURRCY;0SALES_UNIT;0NETVAL_INV;0INV_QTY\n"
                                    "1;202609;RUB;ST;1 000,50;1\n", encoding="utf-8")
    with pytest.raises(ReconcileError, match="нет колонок 0SOLD_TO"):
        read_etalon(tmp_path / "e.csv", spec)
    df = read_etalon(tmp_path / "e.csv", spec, {"Клиент": "0SOLD_TO"})
    assert df.iloc[0]["0NETVAL_INV"] == 1000.5
    (tmp_path / "bad.csv").write_text("0SOLD_TO;0CALMONTH;0DOC_CURRCY;0SALES_UNIT;0NETVAL_INV;0INV_QTY\n"
                                      "1;202609;RUB;ST;много;1\n", encoding="utf-8")
    with pytest.raises(ReconcileError, match="не число"):
        read_etalon(tmp_path / "bad.csv", spec)
    with pytest.raises(ReconcileError, match="не найден"):
        read_etalon(tmp_path / "none.csv", spec)


def test_reconcile_yaml(tmp_path):
    assert load_columns(None, "ZSD_A02") == ({}, 0.01)
    (tmp_path / "reconcile.yaml").write_text("ZSD_A02:\n  columns: {Клиент: 0SOLD_TO}\n  tolerance: 0.5\n",
                                             encoding="utf-8")
    assert load_columns(tmp_path, "ZSD_A02") == ({"Клиент": "0SOLD_TO"}, 0.5)
    (tmp_path / "reconcile.yaml").write_text("ZSD_A02: [", encoding="utf-8")
    with pytest.raises(ReconcileError, match="не читается"):
        load_columns(tmp_path, "ZSD_A02")


def test_excel(spec, tmp_path, model):
    etalon = frame([["1", "202609", "RUB", "ST", 100.0, 1.0]])
    actual = frame([["1", "202609", "RUB", "ST", 90.0, 1.0]])
    detail, totals = compare(etalon, actual, spec)
    lower = ReconcileResult("ZSD_A01", "e1.csv", detail.assign(status=MATCH), totals, spec)
    result = ReconcileResult("ZSD_A02", "e.csv", detail, totals, spec, lower=[lower])
    assert not result.ok and lower.ok
    path = write_excel(result, tmp_path / "r.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Итог", "Расхождения"]
    cells = [c.value for row in wb["Итог"].iter_rows() for c in row]
    assert "Нижний слой ZSD_A01: совпал" in cells
    assert wb["Расхождения"]["K2"].value == DIFF
    assert upstream_adsos(model, "ZSD_A02") == ["ZSD_A01"]
