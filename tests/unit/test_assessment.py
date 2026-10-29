"""Уровни автоматизации, трудоёмкость и Excel-отчёты (ПИ-4, ФП-2.3, ФП-3)."""

import pytest
from conftest import DEMO
from openpyxl import load_workbook

from dbtransfer.analysis.assessment import assess, assess_transformation, load_effort_norms
from dbtransfer.analysis.excel import write_complexity, write_import_check
from dbtransfer.analysis.formulas import unsupported_functions
from dbtransfer.analysis.graph import ScopeError, scope_from
from dbtransfer.importers.bw_csv import import_bw_export
from dbtransfer.model import ObjectRef, ObjectType, Rule, RuleKind, Transformation


@pytest.fixture(scope="module")
def model():
    return import_bw_export(DEMO / "bw_export")


@pytest.fixture(scope="module")
def report(model):
    return assess(model)


def test_levels_match_expected_md(model, report):
    levels = {a.name: a.level for a in report.objects}
    tr = {t.target.name: t.id for t in model.transformations.values()}
    assert levels[tr["ZSD_A01"]] == 1 and levels[tr["ZSD_A02"]] == 1
    assert levels[tr["ZSD_A03"]] == 3 and levels[tr["ZSD_A04"]] == 2
    assert levels["ZSD_CP01"] == 2 and levels["ZOLD_DSO"] == 2
    assert levels["ZSD_Q001"] == 3 and levels["ZMAT_PRODUCT_TREE"] == 3 and levels["ZSD_PC01"] == 3
    assert all(a.level == 1 for a in report.objects if a.type.startswith(("InfoObject", "ADSO", "DTP")))
    assert report.by_level == {1: 26, 2: 3, 3: 4}
    reason = next(a.reason for a in report.objects if a.name == tr["ZSD_A03"])
    assert "конечная рутина ABAP, 10 строк" in reason


def test_effort_hours(report):
    norms = load_effort_norms()
    expected = (norms["routine_abap"] + norms["routine_abap_per_10_lines"]  # 10 строк
                + norms["read_md_rule"] + norms["composite_provider"] + norms["classic_dso"]
                + norms["query"] + norms["hierarchy"] + norms["process_chain"])
    assert report.hours == expected


def test_unsupported_formula():
    assert unsupported_functions("IF( NETVAL > 0; NETVAL; 0 )") == []
    assert unsupported_functions("LOG( NETVAL ) + 'FOO(1)'") == ["LOG"]
    ref = ObjectRef(type=ObjectType.ADSO, name="X")
    t = Transformation(id="T", source=ref, target=ref,
                       rules=[Rule(rule_id=1, kind=RuleKind.FORMULA, param="LOG( A )", target_fields=["B"])])
    a = assess_transformation(t, load_effort_norms())
    assert a.level == 2 and "LOG" in a.reason


def test_complexity_workbook(tmp_path, model, report):
    path = write_complexity(model, report, tmp_path / "complexity.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Итог", "Объекты", "Рутины", "Не разобрано"]
    assert wb["Объекты"].max_row == len(report.objects) + 1
    assert wb["Рутины"]["F2"].value == 10
    summary = [c.value for row in wb["Итог"].iter_rows() for c in row if c.value]
    assert "Всего" in summary and "ZSALES_FF/FLATFILE" in summary


def test_import_check_workbook(tmp_path, model):
    wb = load_workbook(write_import_check(model, tmp_path / "import_check.xlsx"))
    assert wb.sheetnames == ["Файлы", "Проблемы", "Неактивные"]
    assert wb["Неактивные"]["B2"].value == "ZSD_A99"


def test_scope_from_mart(model):
    scope = scope_from(model, "ZSD_A02")
    assert [r.name for r in scope.objects] == ["ZSALES_FF", "ZSD_A01", "ZSD_A02"]
    assert len(scope.transformations) == 2 and len(scope.dtps) == 2
    assert "0NETVAL_INV" in scope.info_objects and "0COUNTRY" in scope.info_objects


def test_scope_unknown_object(model):
    with pytest.raises(ScopeError):
        scope_from(model, "ZNOPE")
